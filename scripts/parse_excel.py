# -*- coding: utf-8 -*-
"""
Parse the TAM pricing source workbooks into a normalized JSON model consumed by
the dashboard (src/data/scopeData.json).

Source of truth:
  - Scope_of_Work_Priced_QTY.xlsx  -> sections / sub-categories / line items
  - Pricing_Master_Annotated.xlsx  -> supporting "Items & Pricing" library + methodology/gaps
  - Riyadh_Driver_Prices.xlsx      -> ground-transport benchmark reference
  - Domestic Exercise.xlsx         -> domestic travel benchmark reference

Run:  python scripts/parse_excel.py
Output: src/data/scopeData.json  (UTF-8, human readable)
"""
import json
import os
import re
import unicodedata

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
# The workbooks live one level up from the dashboard project (the shared pricing folder)
SRC_DIR = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT = os.path.abspath(os.path.join(HERE, "..", "src", "data", "scopeData.json"))

SECTIONS = ["MARKETING", "EVENT MANAGEMENT", "LOGISTICS", "VIDEO PRODUCTIONS"]
SECTION_LABEL = {
    "MARKETING": "Marketing",
    "EVENT MANAGEMENT": "Event Management",
    "LOGISTICS": "Logistics",
    "VIDEO PRODUCTIONS": "Video Productions",
}


def clean(v):
    """Normalize a cell to a trimmed unicode string (or '')."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    # Replace em/en dashes with a plain hyphen (the em dash reads as "AI" copy)
    s = s.replace("—", "-").replace("–", "-")
    s = unicodedata.normalize("NFC", s)
    return s


def num(v):
    """Return a float for numeric cells, else None."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def slugify(*parts):
    base = "-".join(p for p in parts if p)
    base = base.lower()
    base = re.sub(r"[^a-z0-9]+", "-", base)
    return base.strip("-")[:80]


def parse_scope(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Scope of Work"]
    rows = list(ws.iter_rows(values_only=True))

    sections = []
    current_section = None
    current_subcat = "General"
    seen_ids = {}

    for r in rows:
        # pad row to 13 cols
        cells = list(r) + [None] * (13 - len(r))
        main = clean(cells[0])
        sub = clean(cells[1])
        ctype = clean(cells[2])
        qty = num(cells[3])
        unit_price = num(cells[4])
        source_item = clean(cells[7])
        master_cat = clean(cells[8])
        master_qty_unit = clean(cells[9])
        price_source = clean(cells[10])
        match = clean(cells[11])
        note = clean(cells[12])

        main_upper = main.upper()

        # Section header row
        if main_upper in SECTIONS and not sub:
            current_section = {
                "id": slugify(main),
                "key": main_upper,
                "name": SECTION_LABEL[main_upper],
                "items": [],
            }
            sections.append(current_section)
            current_subcat = "General"
            continue

        # Skip subtotal / grand total / meta rows
        if not current_section:
            continue
        if main_upper.endswith("SUBTOTAL") or main_upper == "GRAND TOTAL":
            continue
        if main_upper in ("HOW TO USE:", "MAPPING SUMMARY") or main_upper.startswith(
            "SCOPE LINES"
        ):
            break  # everything after mapping summary is documentation

        # A Main Item value marks a new sub-category grouping. It may appear on
        # its own row, or share the row with the first item of that group.
        if main:
            current_subcat = main
            if not sub:
                continue  # pure grouping row

        # Item row (has a Sub-Item name)
        if not sub:
            continue

        is_not_in_master = match.upper() == "NOT IN MASTER" or "not in Pricing Master" in note or "not in Pricing Master" in clean(
            cells[6]
        )
        priced = unit_price is not None and unit_price > 0 and not is_not_in_master
        if not priced:
            unit_price = None

        item_id = slugify(current_section["key"], current_subcat, sub)
        # de-dup ids
        if item_id in seen_ids:
            seen_ids[item_id] += 1
            item_id = f"{item_id}-{seen_ids[item_id]}"
        else:
            seen_ids[item_id] = 0

        current_section["items"].append(
            {
                "id": item_id,
                "section": current_section["name"],
                "sectionKey": current_section["key"],
                "subCategory": current_subcat,
                "name": sub,
                "type": ctype or "-",
                "defaultQty": int(qty) if qty else (0 if not priced else 1),
                "unitPrice": unit_price,
                "sourceItem": source_item,
                "masterCategory": master_cat,
                "masterQtyUnit": master_qty_unit,
                "priceSource": price_source,
                "matchStatus": match or ("NOT IN MASTER" if not priced else ""),
                "mappingNote": note,
                "isPriced": priced,
            }
        )

    return sections


def parse_pricing_master(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    library = []
    ws = wb["Items & Pricing"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    for r in rows[1:]:
        cells = list(r) + [None] * (12 - len(r))
        item = clean(cells[3])
        if not item:
            continue
        library.append(
            {
                "type": clean(cells[1]),
                "domain": clean(cells[2]),
                "item": item,
                "quantity": clean(cells[4]),
                "price": num(cells[5]),
                "notes": clean(cells[6]),
                "priceSource": clean(cells[7]),
                "scalingDriver": clean(cells[8]),
                "basis": clean(cells[9]),
                "status": clean(cells[10]),
                "reviewNote": clean(cells[11]),
            }
        )

    gaps = []
    if "Gaps & Data Issues" in wb.sheetnames:
        gw = wb["Gaps & Data Issues"]
        for r in gw.iter_rows(values_only=True):
            vals = [clean(c) for c in r]
            if any(vals):
                gaps.append([v for v in vals])

    methodology = []
    if "Methodology" in wb.sheetnames:
        mw = wb["Methodology"]
        for r in mw.iter_rows(values_only=True):
            vals = [clean(c) for c in r]
            if any(vals):
                methodology.append([v for v in vals])

    return {"library": library, "gaps": gaps, "methodology": methodology}


def parse_reference(driver_path, domestic_path):
    ref = {"ground_transport": [], "domestic_flights": [], "notes": []}
    # Riyadh driver summary
    try:
        wb = openpyxl.load_workbook(driver_path, data_only=True)
        ws = wb["Summary"]
        for r in ws.iter_rows(values_only=True):
            vals = [clean(c) for c in r]
            if any(vals):
                ref["ground_transport"].append(vals)
    except Exception as e:
        ref["notes"].append(f"driver parse error: {e}")

    # Domestic flights
    try:
        wb = openpyxl.load_workbook(domestic_path, data_only=True)
        ws = wb["Domestic"]
        rows = list(ws.iter_rows(values_only=True))
        headers = [clean(c) for c in rows[0]]
        for r in rows[1:]:
            vals = [clean(c) for c in r]
            if any(vals):
                ref["domestic_flights"].append(dict(zip(headers, vals)))
    except Exception as e:
        ref["notes"].append(f"domestic parse error: {e}")

    return ref


def main():
    scope_path = os.path.join(SRC_DIR, "Scope_of_Work_Priced_QTY.xlsx")
    master_path = os.path.join(SRC_DIR, "Pricing_Master_Annotated.xlsx")
    driver_path = os.path.join(SRC_DIR, "Riyadh_Driver_Prices.xlsx")
    domestic_path = os.path.join(SRC_DIR, "Domestic Exercise.xlsx")

    sections = parse_scope(scope_path)
    master = parse_pricing_master(master_path)
    reference = parse_reference(driver_path, domestic_path)

    total_items = sum(len(s["items"]) for s in sections)
    priced = sum(1 for s in sections for i in s["items"] if i["isPriced"])
    unpriced = total_items - priced

    data = {
        "meta": {
            "project": "Annual Cultural & Education Awards Program",
            "currency": "SAR",
            "feeRate": 0.15,
            "generatedFrom": "Scope_of_Work_Priced_QTY.xlsx",
            "counts": {
                "sections": len(sections),
                "items": total_items,
                "priced": priced,
                "unpriced": unpriced,
            },
        },
        "sections": sections,
        "pricingMaster": master,
        "reference": reference,
    }

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"Wrote {OUT}")
    print(
        f"sections={len(sections)} items={total_items} priced={priced} unpriced={unpriced} "
        f"library={len(master['library'])}"
    )
    for s in sections:
        sp = sum(1 for i in s["items"] if i["isPriced"])
        print(f"  {s['name']:20} items={len(s['items']):3} priced={sp}")


if __name__ == "__main__":
    main()
